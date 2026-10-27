from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
p=Path(__file__).resolve().parents[1]/'public/templates/BridgeMate_Quiz_Template.docx'
p.parent.mkdir(parents=True,exist_ok=True)
d=Document();sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=sec.bottom_margin=Cm(2);sec.left_margin=sec.right_margin=Cm(2.2)
for name in ['Normal','Title','Heading 1','Heading 2']:
 st=d.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0);st._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman')
d.styles['Normal'].font.size=Pt(12);d.styles['Normal'].paragraph_format.space_after=Pt(5)
d.styles['Normal'].paragraph_format.line_spacing=1.1
d.styles['Title'].font.size=Pt(22);d.styles['Heading 1'].font.size=Pt(16)
d.add_paragraph('Mẫu nhập câu hỏi BridgeMate','Title')
d.add_paragraph('Dùng mẫu này để soạn câu hỏi trắc nghiệm và nhập vào Practice & Learn. Hãy thay các ví dụ ở trang sau bằng nội dung của bạn, rồi lưu dưới định dạng .docx.')
d.add_heading('Quy tắc soạn câu hỏi',1)
for s in ['Mỗi câu bắt đầu bằng Câu 1: hoặc Câu 2: và nội dung câu hỏi. Nhập số câu bằng tay, không dùng danh sách đánh số tự động của Word.','Mỗi lựa chọn nằm trong một đoạn riêng, bắt đầu bằng A., B., C. hoặc D. Cần từ 2 đến 4 lựa chọn liên tiếp.','Dòng Đáp án: ghi đúng một chữ A, B, C hoặc D. Đây là đáp án duy nhất của câu hỏi.','Dòng Giải thích: không bắt buộc. Nội dung có thể dài nhiều dòng.','Bản tiếng Anh không bắt buộc. Dùng các nhãn Question EN:, A EN:, B EN:, C EN:, D EN: và Explanation EN:. Nếu bỏ trống, hệ thống dùng nội dung tiếng Việt.','Giữ nguyên các nhãn trên, có thể thay font, in đậm hoặc màu chữ. Không đưa câu hỏi vào ảnh hoặc công thức dạng hình.','Giới hạn mỗi file là 2 MB và 100 câu. Nội dung câu hỏi và giải thích tối đa 2.000 ký tự, mỗi lựa chọn tối đa 600 ký tự.','Sau khi tải file lên, kiểm tra bản xem trước rồi bấm Xác nhận nhập. Hệ thống thêm câu hỏi mới vào bộ riêng của bạn. Các ví dụ trong file cũng được nhập nếu bạn chưa xóa chúng.']:
 d.add_paragraph('• '+s)
d.add_page_break();d.add_heading('Câu hỏi mẫu để thay nội dung',1)
lines=['Câu 1: Khi chưa hiểu nhiệm vụ, bạn nên làm gì?','Question EN: What should you do when you do not understand a task?','A. Hỏi lại để làm rõ','A EN: Ask for clarification','B. Đoán rồi bắt đầu','B EN: Guess and start','C. Bỏ qua nhiệm vụ','C EN: Ignore the task','Đáp án: A','Giải thích: Hỏi lại từ đầu giúp tránh hiểu lầm và làm sai yêu cầu.','Explanation EN: Asking early helps prevent misunderstandings.','','Câu 2: Thời hạn làm việc nhóm cần thống nhất thông tin nào?','A. Màu nền của tài liệu','B. Múi giờ và thời điểm bàn giao','C. Ảnh đại diện của thành viên','D. Số biểu tượng trong tin nhắn','Đáp án: B','Giải thích: Múi giờ và thời điểm bàn giao giúp mọi thành viên hiểu cùng một hạn chót.']
for s in lines:
 para=d.add_paragraph(s)
 if s.startswith('Câu '):
  para.runs[0].bold=True;para.paragraph_format.space_before=Pt(10)
for elem in d.styles.element.xpath('.//w:pBdr'):
 elem.getparent().remove(elem)
for elem in d.styles.element.xpath('.//w:rFonts'):
 for attr in list(elem.attrib):
  if 'theme' in attr.lower():del elem.attrib[attr]
d.core_properties.title='Mẫu nhập câu hỏi BridgeMate';d.core_properties.author='BridgeMate';d.save(p)
print(p)
